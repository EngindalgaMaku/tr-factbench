#!/usr/bin/env python3
"""
95_build_annotator_c_package.py
Gold-480 üçüncü anotatör (C) paketi. Yayım sonrası duyarlılık analizi için; Gold-480 v1.0 dosyalarını yalnızca okur.

Tasarım (ön kayıt: results/annotation/gold480_annotator_c/PREREGISTRATION_TR.md):
  - 48 A/B uyuşmazlık maddesi + A ile B'nin uzlaştığı maddelerden her sınıftan 8 (toplam 32) dolgu maddesi = 80 madde
  - Dolgu maddeleri sabit tohumla (42) seçilir; 80 madde aynı tohumla karıştırılır
  - Anotatöre yalnızca kör kimlik, bağlam, soru ve iddia verilir; A, B, Gold etiketi ve madde türü gizli
  - Anahtar dosyası (kör kimlik → yayın kimliği, madde türü) paketle paylaşılmaz

Çıktılar:
  package/ANNOTATION_C.xlsx, package/LABEL_GUIDE_TR.md  (anotatöre verilir; package/INSTRUCTIONS_TR.md elle yazılmıştır)
  private/blind_key.csv                                   (.gitignore; sonuçlar gelene kadar paylaşılmaz)
  package_manifest.json                                   (SHA-256 değerleri)
"""
import csv
import hashlib
import json
import random
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

TFB = Path(__file__).resolve().parent.parent
HERE = TFB / "results/annotation/gold480_annotator_c"
GOLD = TFB / "data/evaluation/gold_v1.0"
SEED = 42
FILLER_PER_CLASS = 8
LABELS = ["supported", "partially_supported", "contradicted", "unverifiable"]


def rj(p):
    return [json.loads(l) for l in Path(p).read_text(encoding="utf-8-sig").splitlines() if l.strip()]


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


gold = {r["release_example_id"]: r for r in rj(GOLD / "gold/TR-FactBench_controlled480_GOLD_v1.0.jsonl")}
by_claim = {r["claim"]: eid for eid, r in gold.items()}
A = {by_claim[r["claim"]]: r["annotation_label"] for r in rj(GOLD / "annotators/TR-FactBench_ANNOTATOR_A_RAW_COMPLETE_480_v1.0.jsonl")}
B = {by_claim[r["claim"]]: r["annotation_label"] for r in rj(GOLD / "annotators/TR-FactBench_ANNOTATOR_B_RAW_COMPLETE_480_v1.0.jsonl")}
disputed = sorted(r["release_example_id"] for r in rj(GOLD / "adjudication/ADJUDICATION_LOG_48_v1.0.jsonl"))
assert len(disputed) == 48 and all(A[i] != B[i] for i in disputed)

rng = random.Random(SEED)
agreed = sorted(i for i in gold if A[i] == B[i])
fillers = []
for lab in LABELS:
    pool = [i for i in agreed if gold[i]["gold_label"] == lab]
    fillers += rng.sample(pool, FILLER_PER_CLASS)
items = [(i, "uyusmazlik") for i in disputed] + [(i, "dolgu_uzlasi") for i in fillers]
rng.shuffle(items)

(HERE / "package").mkdir(parents=True, exist_ok=True)
(HERE / "private").mkdir(exist_ok=True)

wb = Workbook()
ws = wb.active
ws.title = "Anotasyon"
head = ["Sıra", "Kimlik", "Bağlam", "Soru", "İddia", "Etiket", "Eminlik (1=az, 2=orta, 3=çok)", "Not (isteğe bağlı)"]
ws.append(head)
key_rows = []
for n, (eid, kind) in enumerate(items, 1):
    blind = "C_" + hashlib.sha256(f"{SEED}:{eid}".encode()).hexdigest()[:10]
    g = gold[eid]
    ws.append([n, blind, g["context"], g["question"], g["claim"], "", "", ""])
    key_rows.append({"sira": n, "kor_kimlik": blind, "release_example_id": eid, "madde_turu": kind,
                     "A": A[eid], "B": B[eid], "gold": g["gold_label"]})
widths = [6, 15, 80, 35, 55, 22, 14, 35]
for col, w in zip("ABCDEFGH", widths):
    ws.column_dimensions[col].width = w
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="DDDDDD")
    cell.alignment = Alignment(wrap_text=True, vertical="top")
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "C2"
dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True)
dv.error = "Lütfen listeden bir etiket seçin."
dv.errorTitle = "Geçersiz etiket"
ws.add_data_validation(dv)
dv.add(f"F2:F{len(items) + 1}")
dv2 = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
ws.add_data_validation(dv2)
dv2.add(f"G2:G{len(items) + 1}")
ws2 = wb.create_sheet("Etiketler")
ws2.append(["Etiket", "Kısa tanım (ayrıntı LABEL_GUIDE_TR.md)"])
for lab, d in [("supported", "İddiadaki bütün atomik önermeler bağlam tarafından desteklenir."),
               ("partially_supported", "Çok atomlu iddianın en az bir atomu desteklenir, en az bir atomu desteklenmez veya çelişir."),
               ("contradicted", "Hiçbir atom desteklenmez ve en az bir atom bağlamdaki açık bilgiyle çelişir."),
               ("unverifiable", "Bütün atomlar bağlamda yok ya da karar vermeye yetmez; açık bir çelişki yoktur.")]:
    ws2.append([lab, d])
ws2.column_dimensions["A"].width = 22
ws2.column_dimensions["B"].width = 100
xlsx = HERE / "package" / "ANNOTATION_C.xlsx"
wb.save(xlsx)

shutil.copyfile(TFB / "docs/LABEL_GUIDE_TR.md", HERE / "package" / "LABEL_GUIDE_TR.md")
with open(HERE / "private" / "blind_key.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(key_rows[0]))
    w.writeheader()
    w.writerows(key_rows)

summary = {"tohum": SEED, "madde_sayisi": len(items), "uyusmazlik": 48, "dolgu": len(fillers),
           "dolgu_sinif_dagilimi": {lab: FILLER_PER_CLASS for lab in LABELS},
           "sha256": {p: sha(HERE / p) for p in ["package/ANNOTATION_C.xlsx", "package/LABEL_GUIDE_TR.md", "package/INSTRUCTIONS_TR.md", "private/blind_key.csv"] if (HERE / p).exists()},
           "kaynak_sha256": {"gold": sha(GOLD / "gold/TR-FactBench_controlled480_GOLD_v1.0.jsonl"),
                             "adjudication": sha(GOLD / "adjudication/ADJUDICATION_LOG_48_v1.0.jsonl")}}
(HERE / "package_manifest.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2))
